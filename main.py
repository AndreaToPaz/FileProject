import pandas as pd
from openpyxl import load_workbook
import shutil
import os
from utils import clean_columns
from dotenv import load_dotenv
from tkinter import Tk
from tkinter.filedialog import askopenfilename

# -------------------------
# Load environment variables
# -------------------------
load_dotenv()
template_file = os.getenv("TEMPLATE_FILE", "template.xlsx")
output_folder = os.getenv("OUTPUT_FOLDER", "output")

# -------------------------
# Open file dialog to pick master file
# -------------------------
Tk().withdraw()  # Hide the main tkinter window
master_file = askopenfilename(
    title="Select Master Excel File",
    filetypes=[("Excel files", "*.xlsx *.xls")]
)

if not master_file:
    raise FileNotFoundError("No master file selected. Exiting.")

print(f"Selected master file: {master_file}")

# -------------------------
# Read and clean master file
# -------------------------
df = pd.read_excel(master_file)
df = clean_columns(df)

# Let user pick target column
print("\nAvailable columns:")
for i, col in enumerate(df.columns, start=1):
    print(f"{i}. {col}")

col_index = int(input("Enter the number of the target column to split by: ")) - 1
target_column = df.columns[col_index]

# -------------------------
# Create output folder
# -------------------------
os.makedirs(output_folder, exist_ok=True)

# -------------------------
# Process each fleet/value in target column
# -------------------------
fleets = df[target_column].dropna().unique()

for fleet in fleets:
    fleet_df = df[df[target_column] == fleet]
    safe_fleet = str(fleet).replace("/", "-").replace("\\", "-").replace(":", "-")
    new_file = os.path.join(output_folder, f"{safe_fleet}.xlsx")

    shutil.copy(template_file, new_file)

    wb = load_workbook(new_file)
    ws = wb.active

    for row_index, row in enumerate(fleet_df.values, start=2):
        for col_index, value in enumerate(row, start=1):
            ws.cell(row=row_index, column=col_index).value = value

    wb.save(new_file)

print("All files created successfully using template formatting!")